#!/usr/bin/env python3
r"""爆款评估体系：发布前结构评分 + 发布后数据分位诊断。

用法:
    python evaluate.py pre  --sheet 03-scripts\xx.sheet.json [--transcript 02-transcript\xx.json] [--out 报告.md]
    python evaluate.py post --data 08-data\后台导出.csv [--window 20] [--out 报告.md]

原则（来自 baokuan-silu）: 不承诺爆款；发布前只评"为某个目标设计得如何"，
发布后只与账号自身基线比较，一次只改一个变量。
"""
from __future__ import annotations
import sys as _portable_sys
from pathlib import Path as _PortablePath
_portable_sys.path.insert(0, str(_PortablePath(__file__).resolve().parents[1] / "lib"))
from aicut.portability import portable_path

import argparse
import csv
import importlib.util
import json
import statistics as st
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SKILL_CHECKER = Path(portable_path('user:.agents/skills/duanshipin-hegui-jianji/scripts/check_hegui.py'))

REVERSE_WORDS = ["其实", "别再", "为什么", "竟然", "反而", "不是", "错了", "误区"]
EVIDENCE_WORDS = ["流程", "清单", "步骤", "标准", "记录", "实测", "对比", "检查", "数据", "案例"]


def load_checker():
    if not SKILL_CHECKER.exists():
        return None
    spec = importlib.util.spec_from_file_location("check_hegui", SKILL_CHECKER)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def load_transcript(path: str | None) -> list[dict]:
    if not path:
        return []
    p = Path(path)
    if not p.is_absolute():
        p = ROOT / p
    if not p.exists():
        return []
    data = json.loads(p.read_text(encoding="utf-8"))
    return data.get("segments", [])


def score_pre(sheet: dict, segments: list[dict]) -> tuple[dict, list[str]]:
    notes: list[str] = []
    dims: dict[str, int] = {}

    texts = [s.get("text", "").strip() for s in segments if s.get("text")]
    first = texts[0] if texts else ""
    cover = (sheet.get("cover_text") or "").strip()
    cta = (sheet.get("cta") or "").strip()
    clips = sheet.get("clips", [])
    total = sum(float(c.get("end", 0)) - float(c.get("start", 0)) for c in clips)

    # 1. 钩子强度 20
    s = 0
    if first:
        s += 6 if len(first) <= 12 else 3
        if any(ch.isdigit() for ch in first):
            s += 4
        if any(w in first for w in REVERSE_WORDS):
            s += 4
    if cover:
        s += 6
    dims["钩子强度"] = min(s, 20)
    if dims["钩子强度"] < 14:
        notes.append("钩子偏弱：首句压到 12 字内，并把结果、数字或反差前置")

    # 2. 承诺兑现 15
    s = 0
    if cover:
        s += 5
    if segments and segments[0].get("start", 99) <= 3:
        s += 5
    if cover and texts and any(tok and tok in "".join(texts[:3]) for tok in [cover[:4]]):
        s += 5
    dims["承诺兑现"] = min(s, 15)
    if dims["承诺兑现"] < 10:
        notes.append("承诺兑现慢：前 3 秒就要出现承诺，并在 15 秒内开始兑现")

    # 3. 结构与推进 20
    s = 0
    if 20 <= total <= 90:
        s += 6
    if len(texts) >= 3:
        s += 7
    if cta:
        s += 7
    dims["结构与推进"] = min(s, 20)
    if dims["结构与推进"] < 14:
        notes.append("结构不完整：补足钩子-展开-落点，并只留一个 CTA")

    # 4. 证据可信 15
    body = "".join(texts)
    s = 8 if any(w in body for w in EVIDENCE_WORDS) or any(ch.isdigit() for ch in body) else 3
    s += 7  # 合规在维度 6 扣分，这里给基础分
    dims["证据可信"] = min(s, 15)
    if s < 12:
        notes.append("证据不足：加入流程、清单或可核验的数字")

    # 5. 情绪与节奏 15
    s = 0
    if len(clips) >= 2:
        s += 7
    per_clip = total / max(len(clips), 1)
    if 2 <= per_clip <= 6:
        s += 8
    else:
        s += 4
    dims["情绪与节奏"] = min(s, 15)
    if dims["情绪与节奏"] < 11:
        notes.append("节奏单一：镜头平均长度控制在 2-6 秒，并让节奏有快慢变化")

    # 6. 落点与合规 15
    s = 8 if cta else 0
    dims["落点与合规"] = s  # 合规部分在 cmd_pre 补齐
    if not cta:
        notes.append("缺少落点：结尾只给一个明确行动指令")

    return dims, notes


def grade(total: int) -> str:
    if total >= 85:
        return "优：可以进入拍摄"
    if total >= 70:
        return "良：补齐短板即可拍"
    if total >= 55:
        return "待改：至少改一项再拍"
    return "需重做：建议换结构或换选题"


def cmd_pre(args) -> int:
    sheet = json.loads(Path(args.sheet).read_text(encoding="utf-8"))
    segments = load_transcript(args.transcript or sheet.get("subtitles_json"))
    dims, notes = score_pre(sheet, segments)

    checker = load_checker()
    compliance = []
    if checker:
        text = "\n".join(s.get("text", "") for s in segments) + "\n" + (sheet.get("cover_text") or "") + "\n" + (sheet.get("cta") or "")
        for h in checker.scan(text):
            if h["方式"] == "原文":  # 只报原文命中，避免变体重复
                compliance.append(f"[{h['类别']}] {h['命中']}")

    cta_present = bool((sheet.get("cta") or "").strip())
    dims["落点与合规"] = (8 if cta_present else 0) + (7 if not compliance else max(0, 7 - 2 * len(compliance)))
    total = sum(dims.values())

    full_map = {"钩子强度": 20, "承诺兑现": 15, "结构与推进": 20, "证据可信": 15, "情绪与节奏": 15, "落点与合规": 15}
    SUGGEST = {
        "钩子强度": "钩子偏弱：首句压到 12 字内，把结果、数字或反差前置",
        "承诺兑现": "承诺兑现慢：前 3 秒出现承诺，15 秒内开始兑现",
        "结构与推进": "结构不完整：补足钩子-展开-落点，只留一个 CTA",
        "证据可信": "证据不足：加入流程、清单或可核验的数字",
        "情绪与节奏": "节奏单一：镜头平均长度 2-6 秒，并做出快慢变化",
        "落点与合规": "落点或合规有问题：结尾只给一个行动指令，并先改写风险表达",
    }
    _have = {n.split("：")[0] for n in notes}
    for _k, _v in dims.items():
        _prefix = SUGGEST[_k].split("：")[0]
        if _v < full_map[_k] * 0.75 and SUGGEST[_k] not in notes and _prefix not in _have:
            notes.append(SUGGEST[_k])
            _have.add(_prefix)

    lines = ["# 发布前评估", "", f"项目：{sheet.get('project', '(未命名)')}", f"总分：**{total}/100** — {grade(total)}", "", "## 维度得分", "",
             "| 维度 | 得分 | 满分 |", "| --- | --- | --- |"]
    full = {"钩子强度": 20, "承诺兑现": 15, "结构与推进": 20, "证据可信": 15, "情绪与节奏": 15, "落点与合规": 15}
    for k, v in dims.items():
        lines.append(f"| {k} | {v} | {full.get(k, '-')} |")

    lines += ["", "## 观察到的短板", ""]
    if notes:
        lines += [f"- {n}" for n in notes]
    else:
        lines.append("- 未发现明显短板")

    lines += ["", "## 合规检查", ""]
    if compliance:
        lines += [f"- {c}" for c in compliance]
        lines.append("- 建议先改写再拍摄")
    else:
        lines.append("- 未命中已知风险表达（不代表平台必然通过）")

    if notes:
        lines += ["", "## 下一步：只改一个变量", "", f"- 本轮建议只改：**{notes[0].split('：')[0]}**", "- 其余部分保持不变，发布后与自身基线对比"]
    lines += ["", "> 本评分只衡量“为某个目标设计得如何”，不代表平台流量结果。", ""]

    report = "\n".join(lines)
    print(report)
    out = Path(args.out) if args.out else ROOT / "08-data" / f"评估-{sheet.get('project', '未命名')}.md"
    if not out.is_absolute():
        out = ROOT / out
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(report, encoding="utf-8")
    print(f"报告已写入 {out}")
    return 0


def read_rows(path: Path) -> list[dict]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(c) if c is not None else "" for c in rows[0]]
        return [dict(zip(header, r)) for r in rows[1:]]
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def num(v) -> float | None:
    try:
        return float(str(v).replace("%", "").replace(",", "").strip())
    except Exception:
        return None


def cmd_post(args) -> int:
    path = Path(args.data)
    if not path.is_absolute():
        path = ROOT / path
    rows = read_rows(path)
    if not rows:
        print("没有读到数据")
        return 1
    cols = {k: next((c for c in rows[0] if k in c), None) for k in ["播放", "完播", "点赞", "收藏", "评论", "涨粉", "标题", "作品"]}
    window = rows[-args.window:] if len(rows) > args.window else rows

    def med(key):
        c = cols.get(key)
        if not c:
            return None
        vals = [v for v in (num(r.get(c)) for r in window) if v is not None]
        return st.median(vals) if vals else None

    baseline = {k: med(k) for k in ["播放", "完播", "点赞", "收藏", "评论", "涨粉"]}
    lines = ["# 发布后诊断", "", f"数据源：{path.name}｜基线窗口：最近 {len(window)} 条", "", "## 自身基线（中位数）", "",
             "| 指标 | 中位数 |", "| --- | --- |"]
    for k, v in baseline.items():
        lines.append(f"| {k} | {'-' if v is None else round(v, 2)} |")

    lines += ["", "## 逐条诊断", ""]
    for r in window:
        title = str(r.get(cols.get("标题") or cols.get("作品") or "", ""))[:26]
        play = num(r.get(cols.get("播放") or ""))
        finish = num(r.get(cols.get("完播") or ""))
        tags = []
        if finish is not None and baseline["完播"] and finish < baseline["完播"] * 0.8:
            tags.append("开头承诺没兑现")
        if play is not None and baseline["播放"] and play < baseline["播放"] * 0.5:
            tags.append("分发不足：换钩子或封面")
        if play is not None and baseline["播放"] and play > baseline["播放"] * 1.3 and finish is not None and baseline["完播"] and finish < baseline["完播"]:
            tags.append("播放好但留存差：选题对、结构要改")
        lines.append(f"- **{title}**" + ("：" + "；".join(tags) if tags else "：未见明显异常"))

    lines += ["", "## 下一步（每次只改一个变量）", "", "| 视频 | 假设 | 只改的变量 | 成功判据 | 观察 | 下一步 |",
              "| --- | --- | --- | --- | --- | --- |", "|  |  |  |  |  |  |", "",
              "> 仅与自身基线比较；平台指标口径会变，以官方后台为准。", ""]
    report = "\n".join(lines)
    print(report)
    if args.out:
        Path(args.out).write_text(report, encoding="utf-8")
        print(f"报告已写入 {args.out}")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description="爆款评估体系")
    sub = ap.add_subparsers(dest="mode", required=True)
    p1 = sub.add_parser("pre"); p1.add_argument("--sheet", required=True); p1.add_argument("--transcript"); p1.add_argument("--out"); p1.set_defaults(func=cmd_pre)
    p2 = sub.add_parser("post"); p2.add_argument("--data", required=True); p2.add_argument("--window", type=int, default=20); p2.add_argument("--out"); p2.set_defaults(func=cmd_post)
    args = ap.parse_args()
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
