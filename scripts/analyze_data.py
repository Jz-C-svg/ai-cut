"""抖音后台数据解析与复盘（AI cut / 阶段：复盘层）。

用法:
    python analyze_data.py 08-data/<导出文件>.csv|xlsx [--name 复盘-2026-09-10] [--window 20]

按 baokuan-silu 的要求：以自己账号近 10-20 条可比作品的中位数为基线，
只输出「观察 / 推断 / 下一次只改一个变量」，不承诺爆款。
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import statistics as st
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent

# 后台导出表头不固定，用关键词做映射
COLUMN_HINTS = {
    "title": ["作品", "标题", "视频"],
    "publish_time": ["发布时间", "日期", "时间"],
    "views": ["播放", "曝光"],
    "finish_rate": ["完播"],
    "avg_watch": ["平均播放", "平均观看"],
    "likes": ["点赞"],
    "comments": ["评论"],
    "shares": ["分享", "转发"],
    "favorites": ["收藏"],
    "fans": ["涨粉", "粉丝增量", "新增粉丝"],
    "followers_rate": ["主页访问", "转化"],
    "duration": ["时长"],
}


def read_rows(path: Path) -> list[dict]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        try:
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return []
            header = [str(c) if c is not None else "" for c in rows[0]]
            return [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]
        finally:
            wb.close()
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def map_columns(row: dict) -> dict:
    """把实际表头映射到标准字段。"""
    mapping = {}
    for field, hints in COLUMN_HINTS.items():
        for col in row:
            if any(h in str(col) for h in hints):
                mapping[field] = col
                break
    return mapping


def to_number(value) -> float | None:
    if value is None:
        return None
    text = str(value).strip().replace(",", "").replace("%", "")
    try:
        return float(text)
    except ValueError:
        return None


def median_of(rows: list[dict], col: str | None) -> float | None:
    if not col:
        return None
    values = [v for v in (to_number(r.get(col)) for r in rows) if v is not None]
    return st.median(values) if values else None


def main() -> int:
    ap = argparse.ArgumentParser(description="AI cut 数据复盘")
    ap.add_argument("file")
    ap.add_argument("--name")
    ap.add_argument("--window", type=int, default=20, help="基线窗口（最近 N 条）")
    args = ap.parse_args()
    if args.window <= 0:
        ap.error("基线窗口必须大于 0")

    path = Path(args.file)
    if not path.is_absolute():
        path = ROOT / path
    rows = read_rows(path)
    if not rows:
        print("没有读到数据行")
        return 1

    mapping = map_columns(rows[0])
    window = rows[-args.window:] if len(rows) > args.window else rows
    name = args.name or f"复盘-{datetime.now():%Y%m%d-%H%M}"

    def med(field: str):
        return median_of(window, mapping.get(field))

    baseline = {
        "样本数": len(window),
        "播放中位数": med("views"),
        "完播率中位数": med("finish_rate"),
        "平均观看中位数": med("avg_watch"),
        "点赞中位数": med("likes"),
        "评论中位数": med("comments"),
        "收藏中位数": med("favorites"),
        "涨粉中位数": med("fans"),
    }

    # 逐条相对基线诊断（四象限）
    findings = []
    for r in window:
        title = str(r.get(mapping.get("title", ""), ""))[:28] or "(无标题)"
        views = to_number(r.get(mapping.get("views", "")))
        finish = to_number(r.get(mapping.get("finish_rate", "")))
        likes = to_number(r.get(mapping.get("likes", "")))
        favorites = to_number(r.get(mapping.get("favorites", "")))
        comments = to_number(r.get(mapping.get("comments", "")))
        notes = []
        if finish is not None and baseline["完播率中位数"] and finish < baseline["完播率中位数"] * 0.8:
            notes.append("完播率低于本窗口中位数的 80%；原因未由这组数据确定")
        if views is not None and baseline["播放中位数"] and views < baseline["播放中位数"] * 0.5:
            notes.append("播放量低于本窗口中位数的 50%；需结合曝光量和发布时间判断原因")
        if favorites is not None and baseline["收藏中位数"] and likes and favorites < likes * 0.05:
            notes.append("收藏/点赞低于 5%；只表示互动比例偏低，不直接代表内容价值")
        if comments is not None and baseline["评论中位数"] and comments < baseline["评论中位数"] * 0.5:
            notes.append("评论数低于本窗口中位数的 50%；不能单独归因为缺少观点或提问")
        if notes:
            findings.append((title, notes))

    out_dir = ROOT / "08-data"
    out_dir.mkdir(parents=True, exist_ok=True)
    md = [f"# {name}", "", f"- 数据来源：`{path.name}`", f"- 基线窗口：最近 {len(window)} 条", ""]
    md += ["## 基线（中位数）", "", "| 指标 | 数值 |", "| --- | --- |"]
    for k, v in baseline.items():
        md.append(f"| {k} | {'-' if v is None else round(v, 2)} |")
    md += ["", "## 观察到的问题", ""]
    if findings:
        for title, notes in findings:
            md.append(f"- **{title}**：" + "；".join(notes))
    else:
        md.append("- 未发现明显低于基线的作品")
    md += ["", "## 下一步（每次只改一个变量）", "",
           "| 视频 | 假设 | 只改的变量 | 成功判据 | 观察 | 下一步 |",
           "| --- | --- | --- | --- | --- | --- |",
           "|  |  |  |  |  |  |", "",
           "> 提醒：以上为相对自身基线的观察，不是平台保证；对外投放与对标数据需按当前平台规则核验。", ""]

    source_sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
    md += [f"- 源文件 SHA-256：`{source_sha256}`",
           f"- 窗口取导出文件末尾 {len(window)} 行；请确认导出行序对应目标时间范围。", ""]
    report = out_dir / f"{name}.md"
    report.write_text("\n".join(md), encoding="utf-8")
    (out_dir / f"{name}.json").write_text(
        json.dumps({"baseline": baseline, "findings": findings, "mapping": mapping,
                    "source_sha256": source_sha256, "window_rows": len(window),
                    "window_order": "last_rows_as_provided_by_export"},
                   ensure_ascii=False, indent=2), encoding="utf-8")
    print("\n".join(md[:24]))
    print(f"\n报告: {report}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
